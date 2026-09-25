from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib import messages
from django.http import HttpResponseForbidden, HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from ..models import Cliente
from ..forms import ClienteCreateForm, ClienteUpdateForm


# ====================== CRUD CLIENTES ======================
@login_required
@permission_required('venta.view_cliente', raise_exception=True)
def consulta_clientes(request):
    if not (request.user.is_superuser or
            request.user.groups.filter(name='grp_cliente').exists() or
            request.user.has_perm('venta.view_cliente')):
        return HttpResponseForbidden('No tiene permisos para ingresar aquí')

    clientes = Cliente.objects.all().order_by('ape_nombre')
    context = {
        'clientes': clientes,
        'titulo': 'Lista de Clientes',
        'mensaje': 'Listado completo de clientes'
    }
    return render(request, 'venta/clientes/lista_clientes.html', context)


@login_required
@permission_required('venta.add_cliente', raise_exception=True)
def crear_cliente(request):
    dni_duplicado = False

    if request.method == 'POST':
        form = ClienteCreateForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cliente registrado correctamente')
            return redirect('crear_cliente')
        else:
            if 'id_cliente' in form.errors:
                for error in form.errors['id_cliente']:
                    if str(error) == "DNI_DUPLICADO":
                        dni_duplicado = True
                        form.errors['id_cliente'].clear()
                        break
    else:
        form = ClienteCreateForm()

    context = {
        'form': form,
        'dni_duplicado': dni_duplicado
    }
    return render(request, 'venta/clientes/crear_cliente.html', context)


@login_required
@permission_required('venta.change_cliente', raise_exception=True)
def actualizar_cliente(request):
    cliente = None
    dni_buscado = None
    form = None

    if request.method == 'POST':
        if 'buscar' in request.POST:
            dni_buscado = request.POST.get('dni_busqueda')
            if dni_buscado:
                try:
                    cliente = Cliente.objects.get(id_cliente=dni_buscado)
                    form = ClienteUpdateForm(instance=cliente)
                    messages.success(request, f'Cliente con DNI {dni_buscado} encontrado')
                except Cliente.DoesNotExist:
                    messages.error(request, 'No se encontró Cliente con ese DNI')
            else:
                messages.error(request, 'Por favor ingrese el DNI para buscar')
        elif 'guardar' in request.POST:
            dni_buscado = request.POST.get('dni_busqueda') or request.POST.get('id_cliente')
            if dni_buscado:
                try:
                    cliente = Cliente.objects.get(id_cliente=dni_buscado)
                    form = ClienteUpdateForm(request.POST, instance=cliente)
                    if form.is_valid():
                        form.save()
                        messages.success(request, 'Cliente actualizado correctamente')
                        cliente.refresh_from_db()
                        form = ClienteUpdateForm(instance=cliente)
                    else:
                        messages.error(request, 'Error en los datos del formulario')
                except Cliente.DoesNotExist:
                    messages.error(request, 'Cliente no encontrado')
            else:
                messages.error(request, 'No se puede identificar al cliente para actualizar')
    context = {
        'form': form,
        'dni_buscado': dni_buscado,
        'cliente_encontrado': cliente is not None,
        'cliente': cliente
    }
    return render(request, 'venta/clientes/u_cliente.html', context)


@login_required
@permission_required('venta.delete_cliente', raise_exception=True)
def borrar_cliente(request):
    clientes_encontrados = []
    tipo_busqueda = 'dni'
    termino_busqueda = ''
    total_registros = 0

    if request.method == 'POST':
        if 'consultar' in request.POST:
            tipo_busqueda = request.POST.get('tipo_busqueda', 'dni')
            termino_busqueda = request.POST.get('termino_busqueda', '').strip()

            if termino_busqueda:
                if tipo_busqueda == 'dni':
                    try:
                        cliente = Cliente.objects.get(id_cliente=termino_busqueda)
                        clientes_encontrados = [cliente]
                    except Cliente.DoesNotExist:
                        messages.error(request, 'No se encontró cliente con ese DNI')
                elif tipo_busqueda == 'nombre':
                    clientes_encontrados = Cliente.objects.filter(
                        ape_nombre__icontains=termino_busqueda
                    ).order_by('id_cliente')

                    if not clientes_encontrados:
                        messages.error(request, 'No se encontraron clientes con ese nombre')

                total_registros = len(clientes_encontrados)
                if total_registros > 0:
                    messages.success(request, f'Se encontraron {total_registros} registro(s)')
            else:
                messages.error(request, 'Ingrese un término de búsqueda')

        elif 'eliminar' in request.POST:
            dni_eliminar = request.POST.get('dni_eliminar')
            if dni_eliminar:
                try:
                    cliente = Cliente.objects.get(id_cliente=dni_eliminar)
                    cliente.delete()
                    messages.success(request, f'Cliente con DNI {dni_eliminar} eliminado correctamente')

                    tipo_busqueda = request.POST.get('tipo_busqueda_actual', 'dni')
                    termino_busqueda = request.POST.get('termino_busqueda_actual', '')

                    if termino_busqueda:
                        if tipo_busqueda == 'dni':
                            clientes_encontrados = []
                        elif tipo_busqueda == 'nombre':
                            clientes_encontrados = Cliente.objects.filter(
                                ape_nombre__icontains=termino_busqueda
                            ).order_by('id_cliente')

                        total_registros = len(clientes_encontrados)
                except Cliente.DoesNotExist:
                    messages.error(request, 'Cliente no encontrado')
    context = {
        'clientes_encontrados': clientes_encontrados,
        'tipo_busqueda': tipo_busqueda,
        'termino_busqueda': termino_busqueda,
        'total_registros': total_registros
    }
    return render(request, 'venta/clientes/borrar_cliente.html', context)


# ====================== EXPORTAR CLIENTES ======================
@login_required
def exportar_clientes_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A1:D1')
    ws['A1'] = "LISTA DE CLIENTES - INFO VENTAS"
    ws['A1'].font = Font(bold=True, size=14, color="1976D2")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ['DNI', 'Apellidos y Nombres', 'Fecha Registro', 'Fecha Sistema']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    ws.row_dimensions[3].height = 25

    clientes = Cliente.objects.all().order_by('ape_nombre')
    for row_num, cliente in enumerate(clientes, 4):
        ws.cell(row=row_num, column=1, value=cliente.id_cliente)
        ws.cell(row=row_num, column=2, value=cliente.ape_nombre)
        ws.cell(row=row_num, column=3, value=cliente.fec_registro.strftime('%d/%m/%Y'))
        ws.cell(row=row_num, column=4, value=cliente.fec_sistema.strftime('%d/%m/%Y %H:%M'))

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 22

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="clientes_info_ventas.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar_clientes_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="clientes_info_ventas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(letter),
                            rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'], fontSize=18,
        textColor=colors.HexColor('#1976D2'), alignment=1, spaceAfter=20,
    )
    elements.append(Paragraph("LISTA DE CLIENTES - INFO VENTAS", titulo_style))
    elements.append(Spacer(1, 10))

    data = [['DNI', 'Apellidos y Nombres', 'Fecha Registro', 'Fecha Sistema']]
    clientes = Cliente.objects.all().order_by('ape_nombre')
    for cliente in clientes:
        data.append([
            cliente.id_cliente,
            cliente.ape_nombre,
            cliente.fec_registro.strftime('%d/%m/%Y'),
            cliente.fec_sistema.strftime('%d/%m/%Y %H:%M'),
        ])

    tabla = Table(data, colWidths=[1.2*inch, 3.5*inch, 1.5*inch, 2*inch])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2196F3')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
    ]))

    elements.append(tabla)
    elements.append(Spacer(1, 15))

    total_style = ParagraphStyle('Total', parent=styles['Normal'], fontSize=11,
                                 textColor=colors.HexColor('#155724'))
    elements.append(Paragraph(f"<b>Total de clientes registrados:</b> {len(clientes)}", total_style))

    doc.build(elements)
    return response