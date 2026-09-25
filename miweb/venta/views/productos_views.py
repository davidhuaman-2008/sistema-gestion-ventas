from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib import messages
from django.http import HttpResponseForbidden, HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from ..models import Producto
from ..forms import ProductoCreateForm, ProductoUpdateForm


# ====================== CRUD PRODUCTOS ======================
@login_required
@permission_required('venta.view_producto', raise_exception=True)
def consulta_productos(request):
    if not (request.user.is_superuser or
            request.user.groups.filter(name='grp_producto').exists() or
            request.user.has_perm('venta.view_producto')):
        return HttpResponseForbidden('No tiene permisos para ingresar aquí')

    productos = list(Producto.objects.all().order_by('nom_prod'))
    productos_activos = sum(1 for p in productos if p.estado)
    productos_inactivos = sum(1 for p in productos if not p.estado)

    context = {
        'productos': productos,
        'productos_activos': productos_activos,
        'productos_inactivos': productos_inactivos,
        'titulo': 'Lista de Productos',
        'mensaje': 'Listado completo de productos'
    }
    return render(request, 'venta/productos/lista_productos.html', context)


@login_required
@permission_required('venta.add_producto', raise_exception=True)
def crear_producto(request):
    if request.method == 'POST':
        form = ProductoCreateForm(request.POST)
        if form.is_valid():
            producto = form.save()
            messages.success(request, f'Producto "{producto.nom_prod}" registrado correctamente')
            return redirect('crear_producto')
    else:
        form = ProductoCreateForm()

    context = {
        'form': form,
        'titulo': 'Registrar Nuevo Producto'
    }
    return render(request, 'venta/productos/crear_producto.html', context)


@login_required
@permission_required('venta.change_producto', raise_exception=True)
def actualizar_producto(request):
    producto = None
    codigo_buscado = None
    form = None

    if request.method == 'POST':
        if 'buscar' in request.POST:
            codigo_buscado = request.POST.get('codigo_busqueda')
            if codigo_buscado:
                try:
                    producto = Producto.objects.get(id_producto=codigo_buscado, estado=True)
                    form = ProductoUpdateForm(instance=producto)
                    messages.success(request, f'Producto con código {codigo_buscado} encontrado')
                except Producto.DoesNotExist:
                    messages.error(request, 'No se encontró producto con ese código o está inactivo')
            else:
                messages.error(request, 'Por favor ingrese el código para buscar')
        elif 'guardar' in request.POST:
            codigo_buscado = request.POST.get('codigo_busqueda') or request.POST.get('id_producto')
            if codigo_buscado:
                try:
                    producto = Producto.objects.get(id_producto=codigo_buscado)
                    form = ProductoUpdateForm(request.POST, instance=producto)
                    if form.is_valid():
                        form.save()
                        messages.success(request, 'Producto actualizado correctamente')
                        producto.refresh_from_db()
                        form = ProductoUpdateForm(instance=producto)
                    else:
                        messages.error(request, 'Error en los datos del formulario')
                except Producto.DoesNotExist:
                    messages.error(request, 'Producto no encontrado')
            else:
                messages.error(request, 'No se puede identificar el producto para actualizar')
    context = {
        'form': form,
        'codigo_buscado': codigo_buscado,
        'producto_encontrado': producto is not None,
        'producto': producto
    }
    return render(request, 'venta/productos/u_producto.html', context)


# ====================== ACTIVAR / DESACTIVAR PRODUCTO ======================
@login_required
@permission_required('venta.change_producto', raise_exception=True)
def activar_desactivar_producto(request):
    productos_encontrados = []
    tipo_busqueda = 'codigo'
    termino_busqueda = ''
    total_registros = 0

    if request.method == 'POST':
        if 'consultar' in request.POST:
            tipo_busqueda = request.POST.get('tipo_busqueda', 'codigo')
            termino_busqueda = request.POST.get('termino_busqueda', '').strip()

            if termino_busqueda:
                if tipo_busqueda == 'codigo':
                    try:
                        producto = Producto.objects.get(id_producto=termino_busqueda)
                        productos_encontrados = [producto]
                    except Producto.DoesNotExist:
                        messages.error(request, 'No se encontró producto con ese código')
                elif tipo_busqueda == 'nombre':
                    productos_encontrados = list(Producto.objects.filter(
                        nom_prod__icontains=termino_busqueda
                    ).order_by('nom_prod'))

                    if not productos_encontrados:
                        messages.error(request, 'No se encontraron productos con ese nombre')

                total_registros = len(productos_encontrados)
                if total_registros > 0:
                    messages.success(request, f'Se encontraron {total_registros} registro(s)')
            else:
                messages.error(request, 'Ingrese un término de búsqueda')

        elif 'cambiar_estado' in request.POST:
            codigo_cambiar = request.POST.get('codigo_cambiar')
            if codigo_cambiar:
                try:
                    producto = Producto.objects.get(id_producto=codigo_cambiar)
                    producto.estado = not producto.estado
                    producto.save()
                    estado_texto = "activado" if producto.estado else "desactivado"
                    messages.success(request, f'Producto "{producto.nom_prod}" {estado_texto} correctamente')

                    tipo_busqueda = request.POST.get('tipo_busqueda_actual', 'codigo')
                    termino_busqueda = request.POST.get('termino_busqueda_actual', '')

                    if termino_busqueda:
                        if tipo_busqueda == 'codigo':
                            productos_encontrados = list(Producto.objects.filter(id_producto=termino_busqueda))
                        elif tipo_busqueda == 'nombre':
                            productos_encontrados = list(Producto.objects.filter(
                                nom_prod__icontains=termino_busqueda
                            ).order_by('nom_prod'))

                        total_registros = len(productos_encontrados)
                except Producto.DoesNotExist:
                    messages.error(request, 'Producto no encontrado')

    context = {
        'productos_encontrados': productos_encontrados,
        'tipo_busqueda': tipo_busqueda,
        'termino_busqueda': termino_busqueda,
        'total_registros': total_registros
    }
    return render(request, 'venta/productos/activar_desactivar_producto.html', context)


# ====================== ELIMINAR PRODUCTO ======================
@login_required
@permission_required('venta.delete_producto', raise_exception=True)
def borrar_producto(request):
    productos_encontrados = []
    tipo_busqueda = 'codigo'
    termino_busqueda = ''
    total_registros = 0

    if request.method == 'POST':
        if 'consultar' in request.POST:
            tipo_busqueda = request.POST.get('tipo_busqueda', 'codigo')
            termino_busqueda = request.POST.get('termino_busqueda', '').strip()

            if termino_busqueda:
                if tipo_busqueda == 'codigo':
                    try:
                        producto = Producto.objects.get(id_producto=termino_busqueda)
                        productos_encontrados = [producto]
                    except Producto.DoesNotExist:
                        messages.error(request, 'No se encontró producto con ese código')
                elif tipo_busqueda == 'nombre':
                    productos_encontrados = list(Producto.objects.filter(
                        nom_prod__icontains=termino_busqueda
                    ).order_by('nom_prod'))

                    if not productos_encontrados:
                        messages.error(request, 'No se encontraron productos con ese nombre')

                total_registros = len(productos_encontrados)
                if total_registros > 0:
                    messages.success(request, f'Se encontraron {total_registros} registro(s)')
            else:
                messages.error(request, 'Ingrese un término de búsqueda')

        elif 'eliminar' in request.POST:
            codigo_eliminar = request.POST.get('codigo_eliminar')
            if codigo_eliminar:
                try:
                    producto = Producto.objects.get(id_producto=codigo_eliminar)
                    nombre = producto.nom_prod
                    producto.delete()
                    messages.success(request, f'Producto "{nombre}" eliminado definitivamente')

                    tipo_busqueda = request.POST.get('tipo_busqueda_actual', 'codigo')
                    termino_busqueda = request.POST.get('termino_busqueda_actual', '')

                    if termino_busqueda:
                        if tipo_busqueda == 'codigo':
                            productos_encontrados = []
                        elif tipo_busqueda == 'nombre':
                            productos_encontrados = list(Producto.objects.filter(
                                nom_prod__icontains=termino_busqueda
                            ).order_by('nom_prod'))

                        total_registros = len(productos_encontrados)
                except Producto.DoesNotExist:
                    messages.error(request, 'Producto no encontrado')

    context = {
        'productos_encontrados': productos_encontrados,
        'tipo_busqueda': tipo_busqueda,
        'termino_busqueda': termino_busqueda,
        'total_registros': total_registros
    }
    return render(request, 'venta/productos/borrar_producto.html', context)


# ====================== EXPORTAR PRODUCTOS ======================
@login_required
def exportar_productos_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A1:F1')
    ws['A1'] = "LISTA DE PRODUCTOS - INFO VENTAS"
    ws['A1'].font = Font(bold=True, size=14, color="1976D2")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ['Código', 'Producto', 'Precio', 'Stock', 'Vencimiento', 'Estado']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    ws.row_dimensions[3].height = 25

    productos = Producto.objects.all().order_by('nom_prod')
    for row_num, producto in enumerate(productos, 4):
        ws.cell(row=row_num, column=1, value=producto.id_producto)
        ws.cell(row=row_num, column=2, value=producto.nom_prod)
        ws.cell(row=row_num, column=3, value=float(producto.precio))
        ws.cell(row=row_num, column=4, value=producto.cantidad)
        fecha_venc = producto.fec_vencimiento.strftime('%d/%m/%Y') if producto.fec_vencimiento else '--'
        ws.cell(row=row_num, column=5, value=fecha_venc)
        ws.cell(row=row_num, column=6, value='Activo' if producto.estado else 'Inactivo')

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 15

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="productos_info_ventas.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar_productos_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="productos_info_ventas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(letter),
                            rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'], fontSize=18,
        textColor=colors.HexColor('#1976D2'), alignment=1, spaceAfter=20,
    )
    elements.append(Paragraph("LISTA DE PRODUCTOS - INFO VENTAS", titulo_style))
    elements.append(Spacer(1, 10))

    data = [['Código', 'Producto', 'Precio', 'Stock', 'Vencimiento', 'Estado']]
    productos = Producto.objects.all().order_by('nom_prod')
    for producto in productos:
        data.append([
            str(producto.id_producto),
            producto.nom_prod,
            f"S/. {producto.precio}",
            str(producto.cantidad),
            producto.fec_vencimiento.strftime('%d/%m/%Y') if producto.fec_vencimiento else '--',
            'Activo' if producto.estado else 'Inactivo',
        ])

    tabla = Table(data, colWidths=[0.9*inch, 2.8*inch, 1*inch, 0.8*inch, 1.3*inch, 1*inch])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2196F3')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
    ]))

    elements.append(tabla)
    elements.append(Spacer(1, 15))

    total_style = ParagraphStyle('Total', parent=styles['Normal'], fontSize=11,
                                 textColor=colors.HexColor('#155724'))
    elements.append(Paragraph(f"<b>Total de productos:</b> {len(productos)}", total_style))

    doc.build(elements)
    return response