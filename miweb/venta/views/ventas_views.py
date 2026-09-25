from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib import messages
from django.http import HttpResponseForbidden, HttpResponse
from django.db import transaction
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from ..models import VentaSimple
from ..forms import VentaSimpleForm


# ====================== CRUD VENTAS SIMPLES ======================
@login_required
@permission_required('venta.view_ventasimple', raise_exception=True)
def consulta_ventas_simples(request):
    if not (request.user.is_superuser or
            request.user.groups.filter(name='grp_venta').exists() or
            request.user.has_perm('venta.view_ventasimple')):
        return HttpResponseForbidden('No tiene permisos para ingresar aquí')

    ventas = VentaSimple.objects.all().order_by('-fecha_venta')
    total_ventas = sum(venta.total for venta in ventas if venta.total)
    total_igv = sum(venta.igv for venta in ventas if venta.igv)

    context = {
        'ventas': ventas,
        'total_ventas': total_ventas,
        'total_igv': total_igv,
        'titulo': 'Lista de Ventas Simples',
        'mensaje': 'Listado completo de ventas'
    }
    return render(request, 'venta/ventas_simples/lista_ventas_simples.html', context)


@login_required
@permission_required('venta.add_ventasimple', raise_exception=True)
@transaction.atomic
def crear_venta_simple(request):
    if request.method == 'POST':
        form = VentaSimpleForm(request.POST)
        if form.is_valid():
            venta = form.save(commit=False)
            producto = venta.producto
            venta.precio_unitario = producto.precio
            venta.subtotal = venta.precio_unitario * Decimal(venta.cantidad)
            venta.igv = venta.subtotal * Decimal('0.18')
            venta.total = venta.subtotal + venta.igv
            venta.save()
            producto.cantidad -= venta.cantidad
            producto.save()
            messages.success(request, 'Venta simple registrada correctamente')
            return redirect('lista_ventas_simples')
    else:
        form = VentaSimpleForm()

    context = {
        'form': form,
        'titulo': 'Registrar Venta Simple',
        'mensaje': 'Complete los datos de la venta'
    }
    return render(request, 'venta/ventas_simples/crear_venta_simple.html', context)


# ====================== BUSCAR VENTA PARA MODIFICAR ======================
@login_required
@permission_required('venta.change_ventasimple', raise_exception=True)
def buscar_venta_modificar(request):
    codigo_buscado = None

    if request.method == 'POST':
        codigo_buscado = request.POST.get('codigo_busqueda')
        if codigo_buscado:
            try:
                venta = VentaSimple.objects.get(cod_venta=codigo_buscado)
                if venta.anulado:
                    messages.warning(request, 'Esta venta está ANULADA. No se puede modificar.')
                else:
                    messages.success(request, f'Venta #{codigo_buscado} encontrada. Redirigiendo a edición...')
                    return redirect('actualizar_venta_simple', pk=venta.cod_venta)
            except VentaSimple.DoesNotExist:
                messages.error(request, 'No se encontró venta con ese código')
        else:
            messages.error(request, 'Por favor ingrese el código de la venta')

    context = {
        'codigo_buscado': codigo_buscado,
        'titulo': 'Modificar Venta Simple'
    }
    return render(request, 'venta/ventas_simples/buscar_venta_modificar.html', context)


# ====================== ACTUALIZAR VENTA (CORREGIDO) ======================
@login_required
@permission_required('venta.change_ventasimple', raise_exception=True)
@transaction.atomic
def actualizar_venta_simple(request, pk):
    venta = get_object_or_404(VentaSimple, pk=pk)
    cantidad_original = venta.cantidad

    if request.method == 'POST':
        form = VentaSimpleForm(request.POST, instance=venta)
        if form.is_valid():
            venta_actualizada = form.save(commit=False)
            producto = venta_actualizada.producto

            # Forzar el precio desde el producto real (evita valores incorrectos)
            venta_actualizada.precio_unitario = producto.precio

            # Recalcular subtotal, IGV y total con Decimal
            venta_actualizada.subtotal = producto.precio * Decimal(venta_actualizada.cantidad)
            venta_actualizada.igv = venta_actualizada.subtotal * Decimal('0.18')
            venta_actualizada.total = venta_actualizada.subtotal + venta_actualizada.igv

            venta_actualizada.save()

            # Actualizar stock si cambió la cantidad
            if cantidad_original != venta_actualizada.cantidad:
                producto.cantidad += (cantidad_original - venta_actualizada.cantidad)
                producto.save()

            messages.success(request, 'Venta simple actualizada correctamente')
            return redirect('lista_ventas_simples')
    else:
        form = VentaSimpleForm(instance=venta)

    context = {
        'form': form,
        'venta': venta,
        'titulo': 'Actualizar Venta Simple'
    }
    return render(request, 'venta/ventas_simples/actualizar_venta_simple.html', context)


# ====================== BUSCAR VENTA PARA ELIMINAR ======================
@login_required
@permission_required('venta.delete_ventasimple', raise_exception=True)
def buscar_venta_eliminar(request):
    codigo_buscado = None

    if request.method == 'POST':
        codigo_buscado = request.POST.get('codigo_busqueda')
        if codigo_buscado:
            try:
                venta = VentaSimple.objects.get(cod_venta=codigo_buscado)
                if venta.anulado:
                    messages.warning(request, 'Esta venta ya está ANULADA')
                else:
                    messages.success(request, f'Venta #{codigo_buscado} encontrada. Redirigiendo para anular...')
                    return redirect('eliminar_venta_simple', pk=venta.cod_venta)
            except VentaSimple.DoesNotExist:
                messages.error(request, 'No se encontró venta con ese código')
        else:
            messages.error(request, 'Por favor ingrese el código de la venta')

    context = {
        'codigo_buscado': codigo_buscado,
        'titulo': 'Eliminar/Anular Venta Simple'
    }
    return render(request, 'venta/ventas_simples/buscar_venta_eliminar.html', context)


@login_required
@permission_required('venta.delete_ventasimple', raise_exception=True)
@transaction.atomic
def eliminar_venta_simple(request, pk):
    venta = get_object_or_404(VentaSimple, pk=pk)

    if not venta.anulado:
        if request.method == 'POST':
            producto = venta.producto
            producto.cantidad += venta.cantidad
            producto.save()

            venta.anulado = True
            venta.save()
            messages.success(request, 'Venta anulada correctamente')
            return redirect('lista_ventas_simples')

        return render(request, 'venta/ventas_simples/eliminar_venta_simple.html',
                      {'venta': venta, 'titulo': 'Anular Venta'})
    else:
        messages.warning(request, 'Esta venta ya estaba anulada')
        return render(request, 'venta/ventas_simples/eliminar_venta_simple.html',
                      {'venta': venta, 'titulo': 'Venta Anulada'})


# ====================== EXPORTAR VENTAS ======================
@login_required
def exportar_ventas_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="9C27B0", end_color="9C27B0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A1:J1')
    ws['A1'] = "LISTA DE VENTAS - INFO VENTAS"
    ws['A1'].font = Font(bold=True, size=14, color="7B1FA2")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ['Código', 'Cliente', 'Producto', 'Cantidad', 'Precio Unit.', 'Subtotal', 'IGV', 'Total', 'Fecha', 'Estado']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    ws.row_dimensions[3].height = 25

    ventas = VentaSimple.objects.all().order_by('-fecha_venta')
    for row_num, venta in enumerate(ventas, 4):
        ws.cell(row=row_num, column=1, value=venta.cod_venta)
        ws.cell(row=row_num, column=2, value=venta.cliente.ape_nombre)
        ws.cell(row=row_num, column=3, value=venta.producto.nom_prod)
        ws.cell(row=row_num, column=4, value=venta.cantidad)
        ws.cell(row=row_num, column=5, value=float(venta.precio_unitario))
        ws.cell(row=row_num, column=6, value=float(venta.subtotal))
        ws.cell(row=row_num, column=7, value=float(venta.igv))
        ws.cell(row=row_num, column=8, value=float(venta.total))
        ws.cell(row=row_num, column=9, value=venta.fecha_venta.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row_num, column=10, value='ANULADA' if venta.anulado else 'ACTIVA')

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['I'].width = 20
    ws.column_dimensions['J'].width = 12

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ventas_info_ventas.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar_ventas_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="ventas_info_ventas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=landscape(letter),
                            rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'], fontSize=18,
        textColor=colors.HexColor('#7B1FA2'), alignment=1, spaceAfter=20,
    )
    elements.append(Paragraph("LISTA DE VENTAS - INFO VENTAS", titulo_style))
    elements.append(Spacer(1, 10))

    data = [['Cód.', 'Cliente', 'Producto', 'Cant.', 'Subtotal', 'IGV', 'Total', 'Fecha', 'Estado']]
    ventas = VentaSimple.objects.all().order_by('-fecha_venta')
    total_general = 0
    total_igv_general = 0

    for venta in ventas:
        data.append([
            str(venta.cod_venta),
            venta.cliente.ape_nombre,
            venta.producto.nom_prod,
            str(venta.cantidad),
            f"S/. {venta.subtotal}",
            f"S/. {venta.igv}",
            f"S/. {venta.total}",
            venta.fecha_venta.strftime('%d/%m/%Y'),
            'ANULADA' if venta.anulado else 'ACTIVA',
        ])
        if not venta.anulado:
            total_general += float(venta.total)
            total_igv_general += float(venta.igv)

    tabla = Table(data, colWidths=[0.7*inch, 1.8*inch, 1.8*inch, 0.6*inch, 0.9*inch, 0.8*inch, 0.9*inch, 1*inch, 0.8*inch])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#9C27B0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
    ]))

    elements.append(tabla)
    elements.append(Spacer(1, 15))

    total_style = ParagraphStyle('Total', parent=styles['Normal'], fontSize=11,
                                 textColor=colors.HexColor('#155724'))
    elements.append(Paragraph(f"<b>Total de ventas:</b> {len(ventas)}", total_style))
    elements.append(Paragraph(f"<b>Total vendido:</b> S/. {total_general:.2f}", total_style))
    elements.append(Paragraph(f"<b>IGV recaudado:</b> S/. {total_igv_general:.2f}", total_style))

    doc.build(elements)
    return response